#!/usr/bin/env python3
"""Inspect existing source workbooks and RAR directories; no model execution."""
import argparse
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import sys
import traceback
import zipfile

from acquire import atomic_json, verify, fingerprint


def file_sha(path):
    return fingerprint(path)["sha256"]

ITEMS = (("workbook", "Zhang_Source_Data.xlsx"), ("shanghai_archive", "sh_outputs.rar"), ("bus_archive", "bus_outputs.rar"))


def copy_member(archive, member, destination, outer_sha, progress):
    """CRC-check a full ZIP stream; verify and retain a partial prefix on retry."""
    info = archive.getinfo(member)
    destination.parent.mkdir(parents=True, exist_ok=True)
    record = destination.with_suffix(destination.suffix + ".json")
    expected = {"outer_sha256": outer_sha, "member": member, "bytes": info.file_size, "zip_crc32": info.CRC}
    recover_complete = destination.exists() and not record.exists()
    if destination.exists() and not recover_complete:
        metadata = json.loads(record.read_text()) if record.exists() else {}
        if not all(metadata.get(key) == value for key, value in expected.items()) or metadata.get("sha256") != file_sha(destination):
            raise ValueError("Existing source-inspection cache failed verification; it was preserved: " + str(destination))
        progress(bytes_copied=info.file_size, expected_bytes=info.file_size, reused_existing=True)
        return {**metadata, "reused_existing": True}
    partial = destination if recover_complete else destination.with_suffix(destination.suffix + ".part")
    resume = partial.stat().st_size if partial.exists() else 0
    if resume > info.file_size:
        raise ValueError("Partial source member exceeds its declared size; partial file preserved")
    if shutil.disk_usage(destination.parent).free < info.file_size - resume + 64 * 2**20:
        raise ValueError("Insufficient free filesystem space for this source member and a 64 MiB margin")
    digest, position, last_report = hashlib.sha256(), 0, -1
    with archive.open(member) as source, partial.open("r+b" if partial.exists() else "w+b") as output:
        while True:
            chunk = source.read(1024 * 1024)
            if not chunk:
                break
            existing = min(len(chunk), max(0, resume - position))
            if existing and output.read(existing) != chunk[:existing]:
                raise ValueError("Partial source member prefix differs from the ZIP; partial file preserved")
            output.write(chunk[existing:])
            digest.update(chunk)
            position += len(chunk)
            if position - last_report >= 32 * 2**20 or position == info.file_size:
                progress(bytes_copied=position, expected_bytes=info.file_size, resumed_prefix_bytes=resume)
                last_report = position
    if position != info.file_size:
        raise ValueError("Source member stream ended before its declared size")
    metadata = {**expected, "sha256": digest.hexdigest(), "resumed_prefix_bytes": resume, "reused_existing": False}
    if partial != destination:
        partial.replace(destination)
    atomic_json(record, metadata)
    return metadata


def scalar(value):
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("Nonfinite spreadsheet numeric value")
        return value
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def inspect_workbook(path, report):
    import openpyxl
    report.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as archive:
        if sum(row.file_size for row in archive.infolist()) > 128 * 2**20:
            raise ValueError("Workbook expanded XML exceeds the 128 MiB inspection limit")
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    result, hits = [], []
    keywords = ("pems", "shanghai", "beijing", "mae", "rmse", "wmape", "table 1", "llm-utp", "ours")
    try:
        if cached.sheetnames != formulas.sheetnames:
            raise ValueError("Cached and formula workbook views have different sheets")
        for index, title in enumerate(cached.sheetnames):
            sheet, original = cached[title], formulas[title]
            if sheet.max_row > 100000 or sheet.max_column > 256 or sheet.max_row * sheet.max_column > 2000000:
                raise ValueError("Workbook sheet exceeds bounded inspection dimensions: " + title)
            output = report / ("sheet-%03d" % index)
            count, formula_count, missing_cache, preview = 0, 0, 0, []
            with output.with_suffix(".csv").open("w", newline="") as stream, output.with_suffix(".jsonl").open("w") as cells:
                writer = csv.writer(stream)
                for row_number, (original_row, cached_row) in enumerate(zip(original.iter_rows(), sheet.iter_rows()), 1):
                    writer.writerow([scalar(cell.value) for cell in cached_row])
                    records = []
                    for raw, value in zip(original_row, cached_row):
                        if raw.value is None:
                            continue
                        count += 1
                        is_formula = raw.data_type == "f"
                        formula_count += int(is_formula)
                        missing_cache += int(is_formula and value.value is None)
                        records.append({"coordinate": raw.coordinate, "type": raw.data_type,
                            "source_value": scalar(raw.value), "cached_value": scalar(value.value)})
                    if not records:
                        continue
                    row = {"row": row_number, "cells": records}
                    cells.write(json.dumps(row, ensure_ascii=False, allow_nan=False) + "\n")
                    values = [cell["cached_value"] if cell["cached_value"] is not None else cell["source_value"] for cell in records]
                    display = {"row": row["row"], "values": values[:24]}
                    if len(preview) < 12:
                        preview.append(display)
                    if any(word in " ".join(map(str, values)).lower() for word in keywords):
                        hits.append({"sheet": title, **display})
            result.append({"title": title, "rows": sheet.max_row, "columns": sheet.max_column,
                "nonempty_cells": count, "formula_cells": formula_count, "formulas_without_cached_values": missing_cache,
                "csv": str(output.with_suffix(".csv")), "cells_jsonl": str(output.with_suffix(".jsonl")), "preview": preview})
    finally:
        cached.close()
        formulas.close()
    payload = {"sheets": result, "keyword_rows": hits, "formula_policy": "cached values retained, formulas recorded but never recalculated"}
    atomic_json(report / "workbook_inventory.json", payload)
    return payload


def inspect_rar(path, report):
    import rarfile
    # RarFile parses headers in Python. Do not call open(), read() or extract().
    with rarfile.RarFile(path, errors="strict") as archive:
        rows = [{"name": row.filename, "bytes": row.file_size, "compressed_bytes": row.compress_size,
            "directory": row.isdir(), "password_required": row.needs_password(), "crc32": row.CRC}
            for row in archive.infolist()]
    payload = {"members": rows, "member_count": len(rows), "declared_expanded_bytes": sum(row["bytes"] for row in rows),
        "scope": "directory metadata only; no payload decompression or payload-integrity claim"}
    atomic_json(report, payload)
    return payload


def run(project, code, report):
    project, code, report = project.resolve(), code.resolve(), report.resolve()
    code.relative_to(project)
    report.relative_to(project)
    report.mkdir(parents=True, exist_ok=True)
    progress = {"phase": "verify_source_archive", "completed_items": 0, "failed_items": 0, "total_items": len(ITEMS),
        "items": {key: {"state": "PENDING"} for key, _ in ITEMS}}
    def update(key=None, **values):
        if key:
            progress["items"][key].update(values)
        progress["completed_items"] = sum(row["state"] == "COMPLETED" for row in progress["items"].values())
        progress["failed_items"] = sum(row["state"] == "FAILED" for row in progress["items"].values())
        atomic_json(report / "source_progress.json", progress)
        print("SOURCE_INSPECT completed=%d/%d failed=%d latest=%s" %
              (progress["completed_items"], len(ITEMS), progress["failed_items"], key or progress["phase"]), flush=True)
    try:
        import openpyxl, rarfile
        if (openpyxl.__version__, rarfile.__version__) != ("3.1.5", "4.2"):
            raise ValueError("Source inspection package versions differ from the pinned environment")
        update()
        lock = json.loads((code / "config/sources.lock.json").read_text())
        asset = next(row for row in lock["lfs_files"] if row["path"] == "Source Data/Source Data.zip")
        source = project / "work/reference/snapshots" / lock["reference_commit"] / asset["path"]
        source.resolve().relative_to(project)
        actual = verify(source, asset)
        if not actual:
            raise ValueError("Existing source ZIP is missing or fails its pinned digest")
        cache = project / "work/source_data_inspection" / actual["sha256"]
        cache.resolve().relative_to(project)
        logs = Path(os.environ.get("TRAFFIC_WORKER_LOG_DIR", str(report / "item_logs")))
        logs.resolve().relative_to(project)
        logs.mkdir(parents=True, exist_ok=True)
        results = {}
        with zipfile.ZipFile(source) as archive:
            for key, basename in ITEMS:
                log = logs / (key + ".log")
                update(key, state="COPYING", log=str(log))
                try:
                    members = [name for name in archive.namelist() if name.replace("\\", "/").split("/")[-1] == basename]
                    if len(members) != 1:
                        raise ValueError("Expected exactly one source member named " + basename)
                    destination = cache / basename
                    metadata = copy_member(archive, members[0], destination, actual["sha256"], lambda **values: update(key, **values))
                    update(key, state="INSPECTING")
                    if key == "workbook":
                        inspected = inspect_workbook(destination, report / "workbook")
                    else:
                        inspected = inspect_rar(destination, report / (key + "_inventory.json"))
                    results[key] = {"state": "COMPLETED", "container": str(destination), "container_metadata": metadata, **inspected}
                    log.write_text(json.dumps(results[key], indent=2, ensure_ascii=False) + "\n")
                    update(key, state="COMPLETED")
                except Exception as exc:
                    log.write_text(traceback.format_exc())
                    results[key] = {"state": "FAILED", "error": str(exc), "log": str(log)}
                    update(key, state="FAILED", error=str(exc))
        progress["phase"] = "FAILED" if progress["failed_items"] else "COMPLETED_WITH_REVIEW"
        update()
        summary = {"execution_status": progress["phase"], "scientific_status": "REPRODUCTION_NOT_ESTABLISHED",
            "source_zip_sha256": actual["sha256"], "items": results, "training_experiments_completed": 0,
            "limitations": ["Workbook values and archive names do not establish dataset, strategy or test-origin correspondence",
                "RAR payloads are listed only, not decompressed or verified as original forecasts",
                "Supplementary figure archives are not assumed to contain main Table 1 predictions",
                "Attention tensors are not loaded; no inference or training runs in this inspection"]}
        atomic_json(report / "source_inspection.json", summary)
        lines = ["# Released Source Data inspection", "", "Execution: " + progress["phase"], ""]
        for key, result in results.items():
            lines += ["## " + key, "", "State: " + result["state"], ""]
            if result["state"] == "FAILED":
                lines += [result["error"], ""]
            elif key == "workbook":
                for sheet in result["sheets"]:
                    lines += ["Sheet %s: %d rows, %d columns, %d nonempty cells, %d missing formula caches." %
                        (sheet["title"], sheet["rows"], sheet["columns"], sheet["nonempty_cells"], sheet["formulas_without_cached_values"])]
                lines += ["", "Keyword rows (first 30; all retained in workbook/workbook_inventory.json):", ""]
                lines += [json.dumps(row, ensure_ascii=False) for row in result["keyword_rows"][:30]]
            else:
                lines += ["Members: %d. Declared payload bytes: %d." % (result["member_count"], result["declared_expanded_bytes"]), ""]
                lines += ["%s | %d bytes" % (row["name"], row["bytes"]) for row in result["members"][:40]]
        lines += ["", "## Scope", ""] + ["- " + value for value in summary["limitations"]]
        lines += ["", "Recovery: resubmit source-inspect; copied containers are hash-verified and partial ZIP streams resume.", ""]
        (report / "SUMMARY.md").write_text("\n".join(lines))
        print("\n".join(lines), flush=True)
        return 2 if progress["failed_items"] else 3
    except Exception as exc:
        progress["phase"] = "FAILED"
        update()
        atomic_json(report / "source_inspection.json", {"execution_status": "FAILED", "error": str(exc), "items": {}, "training_experiments_completed": 0})
        (report / "failure_traceback.txt").write_text(traceback.format_exc())
        (report / "SUMMARY.md").write_text("# Source inspection failed\n\n" + str(exc) + "\n")
        print("SOURCE_INSPECTION_FAILED:", str(exc), flush=True)
        return 2


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    for name in ("project", "code", "report"):
        parser.add_argument("--" + name, type=Path, required=True)
    args = parser.parse_args()
    sys.exit(run(args.project, args.code, args.report))
